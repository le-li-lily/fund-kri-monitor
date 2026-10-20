from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

from .kri import MonitorResult


HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
OK_FILL = PatternFill("solid", fgColor="C6EFCE")
BREACH_FILL = PatternFill("solid", fgColor="FFC7CE")
WARN_FILL = PatternFill("solid", fgColor="FFEB9C")
WHITE_FONT = Font(color="FFFFFF", bold=True)
BOLD_FONT = Font(bold=True)
THIN_BORDER = Border(bottom=Side(style="thin", color="D9E2F3"))


def write_excel_report(result: MonitorResult, output_path: str | Path) -> Path:
    output = Path(output_path)
    output.parent.mkdir(parents=True, exist_ok=True)

    wb = Workbook()
    dashboard = wb.active
    dashboard.title = "KRI Summary"

    _write_title_block(dashboard, result)
    _write_dataframe(dashboard, result.kri_summary, start_row=5, start_col=1, table_name="KRISummary")
    _style_kri_summary(dashboard, start_row=5, rows=len(result.kri_summary))

    dq_ws = wb.create_sheet("Data Quality")
    _write_dataframe(dq_ws, result.data_quality, start_row=1, start_col=1, table_name="DataQuality")
    _style_status_columns(dq_ws)

    breach_ws = wb.create_sheet("Breach Log")
    _write_dataframe(breach_ws, result.breach_log, start_row=1, start_col=1, table_name="BreachLog")
    _style_status_columns(breach_ws)

    action_ws = wb.create_sheet("Action Plan")
    _write_dataframe(action_ws, result.action_plan, start_row=1, start_col=1, table_name="ActionPlan")
    _style_status_columns(action_ws)

    exposure_ws = wb.create_sheet("Exposure Summary")
    _write_dataframe(exposure_ws, result.exposure_summary, start_row=1, start_col=1, table_name="ExposureSummary")
    _format_exposure(exposure_ws)

    holdings_ws = wb.create_sheet("Holdings")
    _write_dataframe(holdings_ws, result.holdings, start_row=1, start_col=1, table_name="Holdings")
    _format_holdings(holdings_ws)

    returns_ws = wb.create_sheet("Portfolio Returns")
    _write_dataframe(returns_ws, result.portfolio_returns, start_row=1, start_col=1, table_name="PortfolioReturns")
    _format_returns(returns_ws)

    thresholds_ws = wb.create_sheet("Thresholds")
    _write_dataframe(thresholds_ws, result.thresholds, start_row=1, start_col=1, table_name="Thresholds")

    history_ws = wb.create_sheet("Trend History")
    _write_dataframe(history_ws, result.history, start_row=1, start_col=1, table_name="TrendHistory")
    _style_status_columns(history_ws)

    for ws in wb.worksheets:
        _freeze_and_autofit(ws)

    wb.save(output)
    return output


def _write_title_block(ws, result: MonitorResult) -> None:
    ws["A1"] = "Fund KRI Monitor"
    ws["A1"].font = Font(size=16, bold=True, color="1F4E78")
    ws["A2"] = "As of Date"
    ws["B2"] = result.as_of_date.strftime("%Y-%m-%d")
    ws["A3"] = "NAV"
    ws["B3"] = result.nav
    ws["B3"].number_format = "#,##0"
    ws["D2"] = "Breaches"
    ws["E2"] = int((result.kri_summary["status"] == "BREACH").sum())
    ws["F2"] = "Watch Items"
    ws["G2"] = int((result.kri_summary["status"] == "WATCH").sum())
    ws["D3"] = "Generated"
    ws["E3"] = pd.Timestamp.now().strftime("%Y-%m-%d %H:%M:%S")
    for cell in ["A2", "A3", "D2", "D3", "F2"]:
        ws[cell].font = BOLD_FONT


def _write_dataframe(ws, df: pd.DataFrame, start_row: int, start_col: int, table_name: str) -> None:
    if df.empty:
        df = pd.DataFrame({"message": ["No records"]})

    for col_idx, column in enumerate(df.columns, start=start_col):
        cell = ws.cell(row=start_row, column=col_idx, value=str(column))
        cell.fill = HEADER_FILL
        cell.font = WHITE_FONT
        cell.alignment = Alignment(horizontal="center")

    for row_offset, row in enumerate(df.itertuples(index=False), start=1):
        for col_offset, value in enumerate(row):
            cell = ws.cell(row=start_row + row_offset, column=start_col + col_offset, value=_excel_value(value))
            cell.border = THIN_BORDER
            if isinstance(value, pd.Timestamp):
                cell.number_format = "yyyy-mm-dd"

    end_row = start_row + len(df)
    end_col = start_col + len(df.columns) - 1
    ref = f"{get_column_letter(start_col)}{start_row}:{get_column_letter(end_col)}{end_row}"
    table = Table(displayName=table_name, ref=ref)
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    ws.add_table(table)


def _style_kri_summary(ws, start_row: int, rows: int) -> None:
    headers = [ws.cell(row=start_row, column=col).value for col in range(1, ws.max_column + 1)]
    status_col = headers.index("status") + 1
    value_col = headers.index("value") + 1
    warning_col = headers.index("warning") + 1
    limit_col = headers.index("limit") + 1
    format_col = headers.index("format") + 1

    for row in range(start_row + 1, start_row + rows + 1):
        metric_format = ws.cell(row=row, column=format_col).value
        for col in [value_col, warning_col, limit_col]:
            if metric_format == "percent":
                ws.cell(row=row, column=col).number_format = "0.00%"
            elif metric_format == "multiple":
                ws.cell(row=row, column=col).number_format = "0.00x"
            else:
                ws.cell(row=row, column=col).number_format = "0.00"
        status_cell = ws.cell(row=row, column=status_col)
        status_cell.font = BOLD_FONT
        status_cell.alignment = Alignment(horizontal="center")
        if status_cell.value == "BREACH":
            status_cell.fill = BREACH_FILL
        elif status_cell.value == "WATCH":
            status_cell.fill = WARN_FILL
        else:
            status_cell.fill = OK_FILL

    status_letter = get_column_letter(status_col)
    ws.conditional_formatting.add(
        f"{status_letter}{start_row + 1}:{status_letter}{start_row + rows}",
        CellIsRule(operator="equal", formula=['"BREACH"'], fill=BREACH_FILL),
    )


def _style_status_columns(ws) -> None:
    for col in range(1, ws.max_column + 1):
        if ws.cell(row=1, column=col).value == "status":
            for row in range(2, ws.max_row + 1):
                cell = ws.cell(row=row, column=col)
                if cell.value == "BREACH":
                    cell.fill = BREACH_FILL
                elif cell.value == "WATCH":
                    cell.fill = WARN_FILL
                else:
                    cell.fill = OK_FILL
                cell.font = BOLD_FONT


def _format_exposure(ws) -> None:
    headers = {ws.cell(row=1, column=col).value: col for col in range(1, ws.max_column + 1)}
    for row in range(2, ws.max_row + 1):
        for name in ["market_value", "gross_exposure", "liquid_5d_value"]:
            if name in headers:
                ws.cell(row=row, column=headers[name]).number_format = "#,##0"
        for name in ["weight", "absolute_weight", "gross_exposure_nav", "liquid_5d_weight"]:
            if name in headers:
                ws.cell(row=row, column=headers[name]).number_format = "0.00%"


def _format_holdings(ws) -> None:
    headers = {ws.cell(row=1, column=col).value: col for col in range(1, ws.max_column + 1)}
    for row in range(2, ws.max_row + 1):
        for name in ["market_value", "gross_exposure", "adv_30d", "five_day_liquidity_value", "stressed_liquidity_value"]:
            if name in headers:
                ws.cell(row=row, column=headers[name]).number_format = "#,##0"
        if "weight" in headers:
            ws.cell(row=row, column=headers["weight"]).number_format = "0.00%"
        if "absolute_weight" in headers:
            ws.cell(row=row, column=headers["absolute_weight"]).number_format = "0.00%"
        if "days_to_liquid_by_adv" in headers:
            ws.cell(row=row, column=headers["days_to_liquid_by_adv"]).number_format = "0.0"
        if "liquid_within_5d" in headers and ws.cell(row=row, column=headers["liquid_within_5d"]).value is False:
            for col in range(1, ws.max_column + 1):
                ws.cell(row=row, column=col).fill = WARN_FILL


def _format_returns(ws) -> None:
    headers = {ws.cell(row=1, column=col).value: col for col in range(1, ws.max_column + 1)}
    for row in range(2, ws.max_row + 1):
        if "date" in headers:
            ws.cell(row=row, column=headers["date"]).number_format = "yyyy-mm-dd"
        for name in ["portfolio_return", "drawdown"]:
            if name in headers:
                ws.cell(row=row, column=headers[name]).number_format = "0.00%"
        if "cumulative_nav_index" in headers:
            ws.cell(row=row, column=headers["cumulative_nav_index"]).number_format = "0.0000"


def _freeze_and_autofit(ws) -> None:
    ws.freeze_panes = "A2"
    for col in range(1, ws.max_column + 1):
        letter = get_column_letter(col)
        max_len = max(len(str(ws.cell(row=row, column=col).value or "")) for row in range(1, ws.max_row + 1))
        ws.column_dimensions[letter].width = min(max(max_len + 2, 12), 48)
    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)


def _excel_value(value):
    if pd.isna(value):
        return None
    if value in [float("inf"), float("-inf")]:
        return None
    if isinstance(value, pd.Timestamp):
        return value.to_pydatetime()
    return value
